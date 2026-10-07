"""Reconcile the user's consolidated workbook without changing any workbook."""
import json
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent
NEW = Path(r'C:\Users\danil\Downloads\leadsQA_209_URL_LEADS_completados.xlsx')


def text(value):
    """Compare visible text without changing source values."""
    return str(value or '').strip()


def normalized(value):
    """Ignore accents, capitalization and redundant whitespace in labels."""
    return ' '.join(''.join(c for c in unicodedata.normalize('NFD', text(value).lower()) if not unicodedata.combining(c)).split())


def records(path):
    """Extract real landing rows, including independent hyperlink targets."""
    sheet = openpyxl.load_workbook(path).worksheets[0]
    result = []
    for row in range(2, sheet.max_row + 1):
        url = text(sheet.cell(row, 4).value)
        if not url:
            continue
        cell = sheet.cell(row, 15)
        result.append(dict(row=row, country=text(sheet.cell(row, 2).value), level=text(sheet.cell(row, 3).value),
                           url=url, location=text(sheet.cell(row, 5).value), link=text(cell.value),
                           target=cell.hyperlink.target if cell.hyperlink else None,
                           email=text(sheet.cell(row, 34).value)))
    return result


def key(record):
    """Match full URL, country and form location; level is checked separately."""
    return (record['url'].rstrip('/'), normalized(record['country']), normalized(record['location']))


current = records(NEW)
previous = records(ROOT / 'leadsQA_avance_44_confirmados.xlsx')
by_link = defaultdict(list)
by_key = defaultdict(list)
for record in current:
    by_link[record['link']].append(record)
    by_key[key(record)].append(record)

# Each confirmed link must be attached to its original landing, not its old row number.
checks = []
for old in previous:
    if not old['link']:
        continue
    attached = by_link.get(old['link'], [])
    checks.append(dict(old_row=old['row'], email=old['email'], link=old['link'],
                       attached=[dict(row=r['row'], same_key=key(r) == key(old), same_level=normalized(r['level']) == normalized(old['level']), level=r['level']) for r in attached],
                       expected_rows=[r['row'] for r in by_key.get(key(old), [])],
                       old_key=key(old), old_level=old['level']))

# Compare the remaining pending scope and isolate rows absent from the 94-row source.
previous_keys = {(key(r), normalized(r['level'])) for r in previous}
blank = [r for r in current if not r['link']]
pending_outside_previous = [r for r in blank if (key(r), normalized(r['level'])) not in previous_keys]
duplicates = {link: [r['row'] for r in rows] for link, rows in by_link.items() if link and len(rows) > 1}
# The original 209-row matrix provides an independent baseline for pre-existing links.
original = records(NEW.parent / 'leadsQA.xlsx')
original_by_row = {r['row']: r for r in original}
baseline_links = [r for r in original if r['link']]
baseline_changes = [dict(row=r['row'], old=original_by_row[r['row']], current=r)
                    for r in current if original_by_row[r['row']]['link'] and
                    (r['link'] != original_by_row[r['row']]['link'] or key(r) != key(original_by_row[r['row']])
                     or normalized(r['level']) != normalized(original_by_row[r['row']]['level']))]
confirmed_links = {r['link'] for r in previous if r['link']}
baseline_set = {r['link'] for r in baseline_links}
unknown_links = [r for r in current if r['link'] and r['link'] not in confirmed_links | baseline_set]
assert len(checks) == 44
assert all(len(c['attached']) == 1 and c['attached'][0]['same_key'] for c in checks)
assert not baseline_changes and not unknown_links and not duplicates
print(json.dumps(dict(total=len(current), filled=len(current)-len(blank), blank=len(blank),
                      blank_rows=[r['row'] for r in blank], duplicate_links=duplicates,
                      baseline_count=len(baseline_links), baseline_changes=baseline_changes, unknown_links=unknown_links,
                      verified_44=len(checks), level_differences=[c for c in checks if not c['attached'][0]['same_level']],
                      pending_outside_previous=pending_outside_previous), ensure_ascii=False, indent=2))
