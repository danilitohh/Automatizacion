"""Recover existing test identities for the 60 pending rows without editing Excel."""
import json
import unicodedata
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent

def norm(value):
    """Normalize accents and labels for conservative source matching."""
    return ' '.join(''.join(c for c in unicodedata.normalize('NFD', str(value or '').strip().lower()) if not unicodedata.combining(c)).split())

current = openpyxl.load_workbook(Path(r'C:\Users\danil\Downloads\leadsQA_209_URL_LEADS_completados.xlsx')).worksheets[0]
previous = openpyxl.load_workbook(ROOT / 'leadsQA_avance_44_confirmados.xlsx')
main = previous['Hoja 1']
records = []
# Prefer the exact URL/country/level/location identity from the previous source.
for row in range(2, 211):
    if current.cell(row, 15).value:
        continue
    raw = [current.cell(row, c).value for c in range(1, 16)]
    matches = [r for r in range(2, 96) if all(norm(main.cell(r, c).value) == norm(current.cell(row, c).value) for c in [2,3,4,5])]
    record = dict(row=row, country=raw[1], level=raw[2], url=str(raw[3]).strip(), location=raw[4], identities=[])
    if len(matches) == 1:
        r = matches[0]
        record.update(old_row=r, name=main.cell(r,33).value, email=main.cell(r,34).value, phone=str(main.cell(r,35).value),
                      prior_status=main.cell(r,38).value, prior_error=main.cell(r,45).value)
    # Additional report rows can retain identities omitted from the trimmed main sheet.
    for values in previous['Form Validation resultados'].iter_rows(min_row=2, values_only=True):
        if str(values[0] or '').strip().rstrip('/') == record['url'].rstrip('/'):
            record['identities'].append(dict(country=values[1],name=values[2],email=values[3],phone=values[4],level=values[5],status=values[7],error=values[14]))
    records.append(record)
assert len(records) == 60
(ROOT / 'pending_60.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(records,ensure_ascii=False,indent=2))
