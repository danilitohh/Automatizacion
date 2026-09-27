"""Checks the high-risk URL mapping and preservation rules for the module."""

from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import SecretStr

from backend.app.config.settings import Settings
from backend.app.main import create_app
from backend.app.modules.strapi_hreflangs.countries import TARGETS, build_url
from backend.app.modules.strapi_hreflangs.report import build_report
from backend.app.modules.strapi_hreflangs.sync import (
    _component_write_payload,
    _href_langs,
    _slug_candidates,
)


def test_hreflang_urls_apply_the_country_slug_rules():
    """Localized career and master's slugs follow the source project rules."""

    targets = {target.hreflang: target for target in TARGETS}
    assert build_url(targets["es-co"], "licenciatura-en-derecho") == "https://utel.edu.mx/colombia/carrera-en-derecho"
    assert build_url(targets["es-cl"], "maestria-en-derecho") == "https://utel.edu.mx/chile/magister-en-derecho"
    assert build_url(targets["es-mx"], "carrera-en-linea") == "https://utel.edu.mx/licenciatura-en-linea"


def test_slug_candidates_only_swap_complete_degree_prefixes():
    """Avoids corrupting slugs whose first word only starts similarly."""

    assert _slug_candidates("licenciatura-en-derecho", "es-co") == ["licenciatura-en-derecho", "carrera-en-derecho"]
    assert _slug_candidates("licenciaturas-ejecutivas", "es-co") == ["licenciaturas-ejecutivas"]


def test_payload_preserves_sibling_seo_values_and_media_ids():
    """Updating MultipleHrefLangs does not remove SEO siblings or populated media."""

    attributes = {
        "seo": {
            "metaTitle": "Programa",
            "metaImage": {"data": {"id": 31, "attributes": {"url": "/cover.png"}}},
            "MultipleHrefLangs": [],
        }
    }
    payload = _component_write_payload(attributes, "seo.MultipleHrefLangs", [{"hrefLang": "es-mx"}])
    assert payload == {
        "seo": {
            "metaTitle": "Programa",
            "metaImage": 31,
            "MultipleHrefLangs": [{"hrefLang": "es-mx"}],
        }
    }
    assert attributes["seo"]["MultipleHrefLangs"] == []


def test_existing_hreflangs_reject_invalid_rows_without_dropping_them():
    """Invalid existing component data is reported instead of silently erased."""

    try:
        _href_langs([{"url": "https://utel.edu.mx/"}])
    except ValueError as error:
        assert "entrada inválida" in str(error)
    else:
        raise AssertionError("Expected malformed current data to be rejected")


def test_report_contains_the_run_mode_and_product_results():
    """The generated report is a valid Excel workbook with localized outcomes."""

    content = build_report([{
        "action": "would-update",
        "product_name": "Licenciatura en Derecho",
        "siu_key": "123",
        "locale": "es-MX",
        "slug": "licenciatura-en-derecho",
        "href_langs": [{"hrefLang": "es-mx", "url": "https://utel.edu.mx/licenciatura-en-derecho"}],
        "added": [{"hrefLang": "es-mx", "url": "https://utel.edu.mx/licenciatura-en-derecho"}],
        "unavailable": [],
    }], apply=False)
    workbook = load_workbook(BytesIO(content), read_only=True)
    assert workbook.sheetnames == ["Resultados", "Resumen"]
    assert workbook["Resultados"]["A2"].value == "Se actualizaría"
    assert workbook["Resumen"]["B1"].value == "Simulación"


def test_apply_requires_an_unconsumed_preview_and_never_returns_the_token(tmp_path):
    """The API exposes the Strapi host but rejects writes without a preview."""

    app = create_app(Settings(
        database_path=tmp_path / "hreflangs.db",
        storage_dir=tmp_path / "storage",
        strapi_url="https://api-cms.utel.edu.mx",
        strapi_token=SecretStr("private-test-token"),
    ))
    with TestClient(app) as client:
        config = client.get("/api/strapi/hreflangs/config")
        rejected = client.post("/api/strapi/hreflangs/run", json={
            "locale": "es-mx",
            "apply": True,
            "confirmation": "APLICAR api-cms.utel.edu.mx",
            "expected_host": "api-cms.utel.edu.mx",
            "preview_job_id": "missing-preview",
        })

    assert config.status_code == 200
    assert config.json()["configured"] is True
    assert config.json()["host"] == "api-cms.utel.edu.mx"
    assert "private-test-token" not in config.text
    assert rejected.status_code == 409
    assert not app.state.strapi_hreflang_jobs
