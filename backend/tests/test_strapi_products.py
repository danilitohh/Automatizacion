import io
import asyncio

import httpx
import pytest
from openpyxl import Workbook
from docx import Document

from backend.app.modules.strapi_products.canonical import add_country_to_canonical
from backend.app.modules.strapi_products.country import detect_country_from_filename
from backend.app.modules.strapi_products.description_runner import StrapiDescriptionRunner
from backend.app.modules.strapi_products.pdp_description import extract_content_description, extract_description, extract_long_description
from backend.app.modules.strapi_products.models import ProductRow
from backend.app.modules.strapi_products.runner import StrapiProductRunner
from backend.app.modules.strapi_products.spreadsheet import read_product_rows
from backend.app.services.google_drive_client import GoogleDriveClient, google_document_id
from backend.app.services.strapi_client import StrapiClient, searchable_program_name


COUNTRIES = {"mexico": "mexico", "peru": "peru"}


def test_canonical_is_safe_and_idempotent():
    original = "https://utel.edu.mx/programa?x=1#seo"
    expected = "https://utel.edu.mx/mexico/programa?x=1#seo"
    assert add_country_to_canonical(original, "México", COUNTRIES) == expected
    assert add_country_to_canonical(expected, "México", COUNTRIES) == expected


def test_canonical_rejects_invalid_host_or_empty_value():
    with pytest.raises(ValueError):
        add_country_to_canonical("https://example.com/programa", "México", COUNTRIES)
    with pytest.raises(ValueError):
        add_country_to_canonical("", "México", COUNTRIES)


def test_country_is_detected_from_filename():
    country = detect_country_from_filename("[AR] Nuevos Productos.xlsx")
    assert country.label == "Argentina"
    assert country.locale == "es-AR"
    assert country.slug == "argentina"


def test_country_filename_requires_supported_code():
    with pytest.raises(ValueError, match="codigo de pais"):
        detect_country_from_filename("Nuevos Productos.xlsx")
    with pytest.raises(ValueError, match="no esta configurado"):
        detect_country_from_filename("[ZZ] Nuevos Productos.xlsx")


def test_extracts_only_text_immediately_after_program_title():
    workbook = Document()
    workbook.add_paragraph("Licenciatura en Sistemas", style="Title")
    workbook.add_paragraph("Esta es la descripción original del programa.")
    workbook.add_paragraph("PERFIL PROFESIONAL", style="Heading 1")
    workbook.add_paragraph("Este texto no debe extraerse.")
    output = io.BytesIO()
    workbook.save(output)
    assert extract_description("Licenciatura en Sistemas", "pdp.docx", output.getvalue()) == "Esta es la descripción original del programa."


def test_extracts_explicit_long_description_section():
    document = Document()
    document.add_paragraph("Descripción larga", style="Heading 1")
    document.add_paragraph("Este es el contenido extenso.")
    document.add_paragraph("ASIGNATURAS", style="Heading 1")
    document.add_paragraph("No debe incluirse.")
    output = io.BytesIO(); document.save(output)
    assert extract_long_description("pdp.docx", output.getvalue()) == "Este es el contenido extenso."


def test_extracts_content_description_before_first_subject_block():
    document = Document()
    document.add_paragraph("Asignaturas", style="Heading 1")
    document.add_paragraph("¿Qué materias se estudian?")
    document.add_paragraph("El plan de estudios brinda una formación integral.")
    document.add_paragraph("1° cuatrimestre", style="Heading 2")
    document.add_paragraph("Materia que no debe incluirse")
    output = io.BytesIO(); document.save(output)
    assert extract_content_description("pdp.docx", output.getvalue()) == "¿Qué materias se estudian?\n\nEl plan de estudios brinda una formación integral."


def test_description_runner_updates_descriptions_and_default_layout(tmp_path):
    document = Document()
    document.add_paragraph("Programa A", style="Title")
    document.add_paragraph("Descripcion PDP original.")
    document.add_paragraph("Asignaturas")
    document.add_paragraph("¿Qué materias se estudian?")
    document.add_paragraph("Contenido de asignaturas.")
    document.add_paragraph("1° cuatrimestre")
    output = io.BytesIO(); document.save(output)
    calls = []

    class FakeClient:
        async def find_product(self, *args, **kwargs):
            return {"id": 12, "attributes": {"title": "Programa A", "shortDescription": "old", "longDescription": "old", "seo": {"LinkCanonical": "unchanged"}}}

        async def update_product(self, *args, **kwargs):
            calls.append((args, kwargs))

    source = tmp_path / "programa-a.docx"
    source.write_bytes(output.getvalue())
    results, summary = asyncio.run(StrapiDescriptionRunner(FakeClient(), "Argentina", "es-AR").run([ProductRow("Sheet", 2, "Programa A", str(source))]))
    assert results[0].status == "UPDATED"
    assert summary.updated == 1
    assert calls == [((12, {"shortDescription": "Descripcion PDP original.", "longDescription": "Descripcion PDP original.", "contentDescription": "¿Qué materias se estudian?\n\nContenido de asignaturas.", "customLayoutPDP": "thirdLayout"}), {})]


def test_description_runner_sets_third_layout_and_product_specific_tabs(tmp_path):
    document = Document()
    document.add_paragraph("Programa A", style="Title")
    document.add_paragraph("Descripción PDP.")
    document.add_paragraph("Asignaturas")
    document.add_paragraph("Contenido del plan.")
    document.add_paragraph("1° cuatrimestre")
    output = io.BytesIO(); document.save(output)
    source = tmp_path / "programa-a.docx"
    source.write_bytes(output.getvalue())
    updates = []

    class FakeClient:
        async def find_product(self, *args, **kwargs):
            return {"id": 12}

        async def get_product_tabs_bullet_section(self, identifier, locale):
            assert identifier == 12
            assert locale == "es-AR"
            return {"id": 42, "hideSection": True}

        async def find_bullet_tab_by_strapi_name(self, name, locale):
            assert locale == "es-AR"
            return {"id": {"Perfil ingreso Programa A": 101, "Perfil egreso Programa A": 102, "Empleabilidad Programa A": 103}[name]}

        async def update_product(self, identifier, attributes):
            updates.append((identifier, attributes))

    results, summary = asyncio.run(
        StrapiDescriptionRunner(FakeClient(), "Argentina", "es-AR").run(
            [ProductRow("Sheet", 2, "Programa A", str(source))]
        )
    )
    assert results[0].status == "UPDATED"
    assert summary.updated == 1
    assert updates[0][1]["customLayoutPDP"] == "thirdLayout"
    assert updates[0][1]["tabsBulletSection"] == {
        "id": 42,
        "idForScrolling": "bannerSectionPdp",
        "hideSection": True,
        "tabs": [{"id": 101}, {"id": 102}, {"id": 103}],
    }


def test_description_runner_exports_private_google_doc_with_drive_client():
    document = Document()
    document.add_paragraph("Programa A", style="Title")
    document.add_paragraph("Descripcion PDP original.")
    document.add_paragraph("Asignaturas")
    document.add_paragraph("¿Qué materias se estudian?")
    document.add_paragraph("Contenido de asignaturas.")
    document.add_paragraph("1° cuatrimestre")
    output = io.BytesIO()
    document.save(output)

    class FakeDriveClient:
        is_configured = True
        has_any_configuration = True

        async def export_docx(self, document_id):
            assert document_id == "google-doc-id"
            return output.getvalue()

    class FakeStrapiClient:
        async def find_product(self, *args, **kwargs):
            return {"id": 12, "attributes": {"title": "Programa A"}}

        async def update_product(self, *args, **kwargs):
            return {}

    row = ProductRow("Bloque 1 Staging", 3, "Programa A", "https://docs.google.com/document/d/google-doc-id/edit")
    results, summary = asyncio.run(
        StrapiDescriptionRunner(
            FakeStrapiClient(), "Argentina", "es-AR", google_drive_client=FakeDriveClient()
        ).run([row])
    )
    assert results[0].status == "UPDATED"
    assert results[0].description == "Descripcion PDP original."


def test_description_runner_includes_siu_key_in_live_product_payload(tmp_path):
    document = Document()
    document.add_paragraph("Licenciatura en Educación para la Sustentabilidad", style="Title")
    document.add_paragraph("Descripción PDP original.")
    document.add_paragraph("Asignaturas")
    document.add_paragraph("Contenido del plan.")
    document.add_paragraph("1° cuatrimestre")
    output = io.BytesIO()
    document.save(output)
    source = tmp_path / "educacion-sustentabilidad.docx"
    source.write_bytes(output.getvalue())
    updates = []

    class FakeStrapiClient:
        async def find_product(self, *args, **kwargs):
            return {"id": 1571}

        async def update_product(self, identifier, attributes):
            updates.append((identifier, attributes))

    async def lookup(program):
        return "20261591"

    results, summary = asyncio.run(
        StrapiDescriptionRunner(
            FakeStrapiClient(), "México", "es-MX", siu_key_lookup=lookup
        ).run([ProductRow("México", 2, "Licenciatura en Educación para la Sustentabilidad", str(source))])
    )
    assert results[0].status == "UPDATED"
    assert results[0].siu_key == "20261591"
    assert summary.updated == 1
    assert updates[0][0] == 1571
    assert updates[0][1]["siuKey"] == "20261591"
    assert updates[0][1]["bannerKey"] == "20261591"


def test_description_runner_updates_all_metadata_and_faq_when_subjects_exist(tmp_path):
    document = Document()
    document.add_paragraph("Licenciatura en Sistemas", style="Title")
    document.add_paragraph("Descripción corta del programa.")
    document.add_paragraph("Descripción larga", style="Heading 1")
    document.add_paragraph("Descripción extendida independiente.")
    document.add_paragraph("ASIGNATURAS", style="Heading 1")
    document.add_paragraph("¿Qué materias se estudian?")
    document.add_paragraph("Contenido de asignaturas.")
    document.add_paragraph("1° cuatrimestre", style="Heading 2")
    document.add_paragraph("Introducción a sistemas")
    document.add_heading("12. Preguntas frecuentes (FAQ)", level=2)
    document.add_paragraph("1. ¿Qué aprenderé?")
    document.add_paragraph("Aprenderás sistemas.")
    source = tmp_path / "programa-completo.docx"
    document.save(source)
    updates = []

    class FakeStrapiClient:
        async def find_product(self, *args, **kwargs): return {"id": 1571}
        async def get_product_common_questions(self, *args): return None
        async def find_subject(self, title, locale):
            assert title == "Introducción a sistemas"
            return {"id": 55}
        async def find_experience_option(self, title, locale):
            assert title == "En línea Home2026"
            return {"id": 66}
        async def get_product_metadata_relations(self, identifier): return {}
        async def find_relation_by_title(self, endpoint, title, locale):
            return {"id": 77 if endpoint == "education-levels" else 88}
        async def update_product(self, identifier, attributes): updates.append((identifier, attributes))

    async def siu_lookup(program): return "20261591"

    result, _ = asyncio.run(StrapiDescriptionRunner(
        FakeStrapiClient(), "México", "es-MX", siu_key_lookup=siu_lookup,
    ).run([ProductRow("Sheet", 2, "Licenciatura en Sistemas", str(source))]))

    assert result[0].status == "UPDATED", result[0].message
    assert len(updates) == 1
    payload = updates[0][1]
    assert payload["shortDescription"] == "Descripción corta del programa."
    assert payload["longDescription"] == "Descripción extendida independiente."
    assert payload["modalities"] == [{"id": 66}]
    assert payload["education_level"] == {"id": 77}
    assert payload["form_education_levels"] == [{"id": 88}]
    assert payload["relatedProducts"] == [{"id": 1571}]
    assert payload["siuKey"] == payload["bannerKey"] == "20261591"
    assert payload["subjects"] == [{"id": 55}]
    assert payload["commonQuestions"]["dropdowns"] == [{
        "dropdownTitle": "¿Qué aprenderé?", "richText": "Aprenderás sistemas.", "iconPosition": "Left",
    }]


def test_spreadsheet_reads_program_and_skips_blank_rows():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Programa", "tabs"])
    sheet.append(["Programa A", "x"])
    sheet.append([None, None])
    content = io.BytesIO()
    workbook.save(content)
    rows = read_product_rows(content.getvalue())
    assert [(row.sheet, row.row_number, row.program) for row in rows] == [("Sheet", 2, "Programa A")]


def test_spreadsheet_requires_program_column():
    workbook = Workbook()
    workbook.active.append(["URL"])
    content = io.BytesIO()
    workbook.save(content)
    with pytest.raises(ValueError, match="Programa"):
        read_product_rows(content.getvalue())


def test_spreadsheet_finds_header_after_blank_row_and_strips_excel_apostrophe():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([])
    sheet.append([None, "'Programa", "'Tabs"])
    sheet.append([None, "'Programa A", "'Ok"])
    content = io.BytesIO()
    workbook.save(content)
    rows = read_product_rows(content.getvalue())
    assert rows[0].row_number == 3
    assert rows[0].program == "Programa A"


def test_spreadsheet_prefers_document_hyperlink_over_visible_name():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Programa", "Documento PDP"])
    sheet.append(["Programa A", "Programa A_PDP_Argentina"])
    sheet["B2"].hyperlink = "https://docs.google.com/document/d/example/edit?usp=sharing"
    content = io.BytesIO()
    workbook.save(content)
    rows = read_product_rows(content.getvalue())
    assert rows[0].document == "https://docs.google.com/document/d/example/edit?usp=sharing"


def test_google_document_id_supports_docs_and_drive_links():
    assert google_document_id("https://docs.google.com/document/d/doc-123/edit") == "doc-123"
    assert google_document_id("https://drive.google.com/file/d/file-456/view") == "file-456"
    assert google_document_id("https://drive.google.com/open?id=file-789") == "file-789"
    assert google_document_id("https://example.com/document/d/nope") is None


def test_google_drive_client_refreshes_token_and_exports_docx():
    requests = []

    async def handler(request):
        requests.append(request)
        if request.url.host == "oauth2.googleapis.com":
            return httpx.Response(200, request=request, json={"access_token": "temporary-token", "expires_in": 3600})
        assert request.headers["Authorization"] == "Bearer temporary-token"
        assert request.url.path.endswith("/files/doc-123/export")
        assert request.url.params["mimeType"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        return httpx.Response(200, request=request, content=b"docx-bytes")

    async def run():
        http_client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
        drive = GoogleDriveClient("client-id", "client-secret", "refresh-token", client=http_client)
        try:
            assert await drive.export_docx("doc-123") == b"docx-bytes"
            assert await drive.export_docx("doc-123") == b"docx-bytes"
            assert len(requests) == 3  # one token refresh, two document exports
        finally:
            await http_client.aclose()

    asyncio.run(run())


def test_client_update_product_descriptions_sends_only_description_fields():
    requests = []

    async def handler(request):
        requests.append(request)
        return httpx.Response(200, request=request, json={"data": {"id": 1571}})

    async def run():
        http_client = httpx.AsyncClient(transport=httpx.MockTransport(handler), base_url="https://example.test")
        client = StrapiClient("https://example.test", "secret", client=http_client)
        try:
            result = await client.update_product_descriptions(1571, "Texto corto", "Texto largo")
            assert result["data"]["id"] == 1571
            assert requests[0].method == "PUT"
            assert requests[0].url.path.endswith("/api/products/1571")
            assert requests[0].read() == b'{"data":{"shortDescription":"Texto corto","longDescription":"Texto largo"}}'
        finally:
            await http_client.aclose()

    asyncio.run(run())


def test_searchable_program_name_removes_degree_prefix_only():
    assert searchable_program_name("Licenciatura en Ingeniería Robótica") == "Ingeniería Robótica"
    assert searchable_program_name("Maestría Administración") == "Administración"
    assert searchable_program_name("Doctorado en Educación") == "Educación"
    assert searchable_program_name("Ingeniería Robótica") == "Ingeniería Robótica"


def test_client_falls_back_without_locale_for_available_draft():
    requests = []

    async def handler(request):
        requests.append(dict(request.url.params))
        if "locale" in request.url.params:
            return httpx.Response(200, request=request, json={"data": []})
        return httpx.Response(200, request=request, json={"data": [{"id": 1571, "attributes": {"title": "Licenciatura en Ingeniería en Ciencia de Datos e Inteligencia Analítica", "locale": "es-MX"}}]})

    async def run():
        http_client = httpx.AsyncClient(transport=httpx.MockTransport(handler), base_url="https://example.test")
        client = StrapiClient("https://example.test", "secret", client=http_client)
        try:
            result = await client.find_product("Licenciatura en Ingeniería en Ciencia de Datos e Inteligencia Analítica", "es-AR")
            assert result["id"] == 1571
            assert requests[0]["filters[title][$eq]"] == "Ingeniería en Ciencia de Datos e Inteligencia Analítica"
            assert "locale" not in requests[1]
        finally:
            await http_client.aclose()

    asyncio.run(run())


def test_runner_updates_only_seo_and_supports_id():
    calls = []

    class FakeClient:
        async def find_product(self, *args, **kwargs):
            return {"id": 7, "attributes": {"title": "Programa", "seo": {"id": 9, "LinkCanonical": "https://utel.edu.mx/programa", "MetaTitle": "keep"}}}

        async def update_product(self, *args, **kwargs):
            calls.append((args, kwargs))

    results, _ = asyncio.run(StrapiProductRunner(FakeClient(), "México", "es-MX", COUNTRIES).run([ProductRow("Sheet", 2, "Programa")]))
    assert results[0].status == "UPDATED"
    assert calls == [((7, {"seo": {"id": 9, "LinkCanonical": "https://utel.edu.mx/mexico/programa", "MetaTitle": "keep"}}), {})]


def test_runner_reports_each_completed_product_without_shortening_the_batch():
    events = []

    class FakeClient:
        async def find_product(self, program, *args, **kwargs):
            return {"id": program, "attributes": {"seo": {"LinkCanonical": "https://utel.edu.mx/programa"}}}

        async def update_product(self, *args, **kwargs):
            return {}

    rows = [ProductRow("Sheet", index + 2, f"Programa {index + 1}") for index in range(45)]
    results, summary = asyncio.run(StrapiProductRunner(
        FakeClient(), "México", "es-MX", COUNTRIES,
    ).run(rows, on_result=lambda index, total, result, elapsed: events.append((index, total, result.program, elapsed))))

    assert len(results) == summary.total == len(events) == 45
    assert [event[0] for event in events] == list(range(1, 46))
    assert all(event[1] == 45 and event[3] >= 0 for event in events)
    assert events[-1][2] == "Programa 45"


def test_client_retries_transient_error():
    attempts = 0

    async def handler(request):
        nonlocal attempts
        attempts += 1
        if attempts == 1:
            return httpx.Response(503, request=request)
        return httpx.Response(200, request=request, json={"data": [{"id": 1, "attributes": {}}]})

    transport = httpx.MockTransport(handler)
    http_client = httpx.AsyncClient(transport=transport, base_url="https://example.test")
    client = StrapiClient("https://example.test", "secret", client=http_client)
    try:
        result = asyncio.run(client.find_product("Programa", "es-MX"))
        assert result["id"] == 1
        assert attempts == 2
    finally:
        asyncio.run(http_client.aclose())
