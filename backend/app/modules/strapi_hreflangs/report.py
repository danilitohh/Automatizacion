"""Genera el libro Excel descargable de una simulación o aplicación."""

from __future__ import annotations

import io
import json
from collections import Counter
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


ACTION_LABELS = {"updated": "Actualizado", "would-update": "Se actualizaría", "unchanged": "Sin cambios", "error": "Error"}
COUNTRY_NAMES = {
    "x-default": "Predeterminado", "es-mx": "México", "es-co": "Colombia", "es-pe": "Perú",
    "es-ec": "Ecuador", "es-us": "Estados Unidos", "es-ar": "Argentina", "es-do": "República Dominicana",
    "es-gt": "Guatemala", "es-cl": "Chile", "es-sv": "El Salvador", "es-bo": "Bolivia",
    "es-pa": "Panamá", "es-py": "Paraguay",
}


def build_report(results: list[dict[str, Any]], apply: bool) -> bytes:
    """Construye hojas de resultados y resumen en memoria para descarga."""

    workbook = Workbook()
    detail = workbook.active
    detail.title = "Resultados"
    headers = ["Estado", "Nombre del producto", "siuKey", "Locale fuente", "Slug", "Hreflangs finales", "Agregados", "Locales no disponibles", "Error"]
    detail.append(headers)
    for result in results:
        included = [COUNTRY_NAMES.get(item.get("hrefLang", ""), item.get("hrefLang", "")) for item in result.get("href_langs", [])]
        missing = [COUNTRY_NAMES.get(item.get("hreflang", ""), item.get("hreflang", "")) for item in result.get("unavailable", [])]
        detail.append([
            ACTION_LABELS.get(result["action"], result["action"]), result.get("product_name"), result.get("siu_key"),
            result.get("locale"), result.get("slug"), json.dumps(result.get("href_langs", []), ensure_ascii=False),
            json.dumps(result.get("added", []), ensure_ascii=False), ", ".join(missing), result.get("error", ""),
        ])
    overview = workbook.create_sheet("Resumen")
    counts = Counter(item["action"] for item in results)
    overview.append(["Tipo de ejecución", "Aplicar" if apply else "Simulación"])
    overview.append(["Productos procesados", len(results)])
    for action, label in ACTION_LABELS.items():
        overview.append([label, counts[action]])

    header_fill = PatternFill("solid", fgColor="FF1F4E78")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2" if sheet is detail else "A1"
        sheet.auto_filter.ref = sheet.dimensions if sheet is detail else None
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        if sheet is detail:
            for index, width in enumerate((18, 48, 18, 14, 52, 70, 70, 52, 56), start=1):
                sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                state = row[0]
                state.fill = PatternFill("solid", fgColor={"Error": "FFFCE8E6", "Actualizado": "FFE6F4EA", "Se actualizaría": "FFFFF4CE"}.get(str(state.value), "FFF3F4F6"))
                state.font = Font(bold=True)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
