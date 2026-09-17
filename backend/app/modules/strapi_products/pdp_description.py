from __future__ import annotations

import re
import unicodedata
from io import BytesIO
from pathlib import Path
from urllib.parse import urlparse

from docx import Document
from pypdf import PdfReader


def _normalize(value: str) -> str:
    plain = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    return " ".join(re.sub(r"[^a-z0-9 ]+", " ", plain.casefold()).split())


def _clean(value: str) -> str:
    return " ".join(value.replace("\u00a0", " ").split()).strip()


def _is_heading(text: str) -> bool:
    return len(text) <= 140 and (text.isupper() or re.match(r"^(?:\d+[.)]?|perfil|plan de estudios|competencias|informacion)\b", text, re.I))


def _extract_after_title(program: str, blocks: list[str], source: str) -> str:
    target = _normalize(program)
    title_index = next((index for index, block in enumerate(blocks) if _normalize(block) == target or target in _normalize(block)), None)
    if title_index is None:
        raise ValueError(f"No se encontro el titulo del programa en {source}.")

    for candidate in blocks[title_index + 1:]:
        text = _clean(candidate)
        if not text:
            continue
        if _is_heading(text):
            raise ValueError(f"No se encontro una descripcion debajo del titulo en {source}.")
        return text
    raise ValueError(f"No se encontro una descripcion debajo del titulo en {source}.")


def extract_content_description(filename: str, content: bytes) -> str:
    """Extract text after Asignaturas and before the first cuatrimestre block."""
    extension = Path(urlparse(filename).path or filename).suffix.casefold()
    try:
        if extension == ".docx":
            document = Document(BytesIO(content))
            blocks = [paragraph.text for paragraph in document.paragraphs]
        elif extension == ".pdf":
            reader = PdfReader(BytesIO(content))
            blocks = []
            for page in reader.pages:
                blocks.extend((page.extract_text() or "").splitlines())
        else:
            raise ValueError(f"Formato de Documento PDP no soportado: {filename}.")
    except ValueError:
        raise
    except Exception as error:
        raise ValueError(f"No se pudo abrir el Documento PDP {filename}.") from error
    heading_index = next((index for index, block in enumerate(blocks) if _normalize(block) == "asignaturas"), None)
    if heading_index is None:
        raise ValueError(f"No se encontro la seccion Asignaturas en {filename}.")
    collected: list[str] = []
    for block in blocks[heading_index + 1:]:
        text = _clean(block)
        if not text:
            continue
        if re.match(r"^\d+\s*[°º.]?\s*cuatrimestre\b", _normalize(text), re.I):
            break
        collected.append(text)
    if not collected:
        raise ValueError(f"No se encontro contenido antes de las asignaturas en {filename}.")
    return "\n\n".join(collected)


def extract_description(program: str, filename: str, content: bytes) -> str:
    extension = Path(urlparse(filename).path or filename).suffix.casefold()
    try:
        if extension == ".docx":
            document = Document(BytesIO(content))
            blocks = [paragraph.text for paragraph in document.paragraphs]
            return _extract_after_title(program, blocks, filename)
        if extension == ".pdf":
            reader = PdfReader(BytesIO(content))
            blocks = []
            for page in reader.pages:
                blocks.extend((page.extract_text() or "").splitlines())
            return _extract_after_title(program, blocks, filename)
    except ValueError:
        raise
    except Exception as error:
        raise ValueError(f"No se pudo abrir el Documento PDP {filename}.") from error
    raise ValueError(f"Formato de Documento PDP no soportado: {filename}.")


def extract_long_description(filename: str, content: bytes) -> str | None:
    """Extract an explicitly labeled long-description section when present."""
    extension = Path(urlparse(filename).path or filename).suffix.casefold()
    try:
        if extension == ".docx":
            blocks = [paragraph.text for paragraph in Document(BytesIO(content)).paragraphs]
        elif extension == ".pdf":
            blocks = [line for page in PdfReader(BytesIO(content)).pages for line in (page.extract_text() or "").splitlines()]
        else:
            raise ValueError(f"Formato de Documento PDP no soportado: {filename}.")
    except ValueError:
        raise
    except Exception as error:
        raise ValueError(f"No se pudo abrir el Documento PDP {filename}.") from error

    headings = {"descripcion larga", "descripcion extendida", "descripcion del programa", "acerca del programa"}
    start = next((index for index, block in enumerate(blocks) if _normalize(_clean(block)).rstrip(":") in headings), None)
    if start is None:
        return None

    collected: list[str] = []
    for block in blocks[start + 1:]:
        text = _clean(block)
        if not text:
            continue
        if _is_heading(text):
            break
        collected.append(text)
    if not collected:
        raise ValueError(f"La sección de descripción larga está vacía en {filename}.")
    return "\n\n".join(collected)


def extract_program_durations(filename: str, content: bytes) -> list[str]:
    """Return the distinct month durations from the PDP's Duración line."""
    extension = Path(urlparse(filename).path or filename).suffix.casefold()
    try:
        if extension == ".docx":
            blocks = [paragraph.text for paragraph in Document(BytesIO(content)).paragraphs]
        elif extension == ".pdf":
            blocks = []
            for page in PdfReader(BytesIO(content)).pages:
                blocks.extend((page.extract_text() or "").splitlines())
        else:
            raise ValueError(f"Formato de Documento PDP no soportado: {filename}.")
    except ValueError:
        raise
    except Exception as error:
        raise ValueError(f"No se pudo abrir el Documento PDP {filename}.") from error
    durations: list[str] = []
    for block in blocks:
        if not re.search(r"\bduraci[oó]n\s*:", block, re.I):
            continue
        for months in re.findall(r"\b(\d+)\s*mes(?:es)?\b", block, re.I):
            value = f"{months} meses"
            if value not in durations:
                durations.append(value)
    if len(durations) != 2:
        raise ValueError(
            f"El PDP {filename} debe contener exactamente dos duraciones distintas; "
            f"se encontraron: {', '.join(durations) or 'ninguna'}."
        )
    return sorted(durations, key=lambda value: int(value.split()[0]))


def extract_subjects(filename: str, content: bytes) -> list[str]:
    """Extract individual subjects listed under semester/cuatrimestre blocks."""
    extension = Path(urlparse(filename).path or filename).suffix.casefold()
    try:
        if extension == ".docx":
            document = Document(BytesIO(content))
            blocks = [(p.text, p.style.name if p.style else "") for p in document.paragraphs]
        elif extension == ".pdf":
            blocks = [line for page in PdfReader(BytesIO(content)).pages for line in (page.extract_text() or "").splitlines()]
        else:
            raise ValueError(f"Formato de Documento PDP no soportado: {filename}.")
    except ValueError:
        raise
    except Exception as error:
        raise ValueError(f"No se pudo abrir el Documento PDP {filename}.") from error
    subjects: list[str] = []
    in_plan = False
    for item in blocks:
        text = _clean(item[0] if isinstance(item, tuple) else item)
        style = (item[1] if isinstance(item, tuple) else "").casefold()
        normalized = _normalize(text)
        if re.match(r"^\d+\s*[°º.]?\s*(cuatrimestre|semestre)\b", normalized):
            in_plan = True
            continue
        if not in_plan or not text:
            continue
        if re.match(r"^(?:preguntas frecuentes|faq|que es|qué es|requisitos de admision)\b", normalized) or text.startswith(("¿", "?")) or re.match(r"^\d+\.\s*", text):
            break
        if re.match(r"^(?:plan de estudios|asignaturas|materias|formacion|area)\b", normalized):
            continue
        clean = re.sub(r"^[•·\-–]\s*", "", text).strip()
        # PDP DOCX files commonly use plain paragraphs (Normal style) for
        # subjects, while PDFs and other files use bullets.
        if clean and clean not in subjects and not clean.endswith(":"):
            subjects.append(clean)
    return subjects

