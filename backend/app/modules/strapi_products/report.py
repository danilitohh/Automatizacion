from __future__ import annotations

import io

from openpyxl import Workbook

from .models import ProductResult, ProductSummary


def build_report(results: list[ProductResult], summary: ProductSummary) -> bytes:
    workbook = Workbook()
    detail = workbook.active
    detail.title = "Resultados"
    detail.append(["Hoja", "Fila", "Programa", "Pais", "Estado", "Canonical anterior", "Canonical nuevo", "Mensaje"])
    for result in results:
        detail.append([result.sheet, result.row_number, result.program, result.country, result.status, result.old_canonical, result.new_canonical, result.message])

    overview = workbook.create_sheet("Resumen")
    overview.append(["Metrica", "Valor"])
    for key, value in summary.as_dict().items():
        overview.append([key, value])

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_description_report(results: list[ProductResult], summary: ProductSummary) -> bytes:
    workbook = Workbook()
    detail = workbook.active
    detail.title = "Resultados"
    detail.append(["Hoja", "Fila", "Programa", "Pais", "Estado", "siuKey", "bannerKey", "Descripcion extraida", "Mensaje"])
    for result in results:
        detail.append([result.sheet, result.row_number, result.program, result.country, result.status, result.siu_key, result.banner_key, result.description, result.message])
    overview = workbook.create_sheet("Resumen")
    overview.append(["Metrica", "Valor"])
    for key, value in summary.as_dict().items():
        overview.append([key, value])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
