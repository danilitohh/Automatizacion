from __future__ import annotations

import re
import unicodedata
import io
from pathlib import Path

from openpyxl import load_workbook


def _normalize(value: str) -> str:
    plain = unicodedata.normalize("NFKD", str(value)).encode("ascii", "ignore").decode()
    return " ".join(re.sub(r"[^a-z0-9 ]+", " ", plain.casefold()).split())


class FichasLookup:
    """Resolve each program and country to its short ficha URL."""

    def __init__(self, path: str | Path | bytes) -> None:
        self.path = Path(path) if not isinstance(path, bytes) else None
        self._rows: dict[str, str] = {}
        source = io.BytesIO(path) if isinstance(path, bytes) else self.path
        workbook = load_workbook(source, data_only=False, read_only=True)
        sheet = workbook.active
        for row in sheet.iter_rows(min_row=2, values_only=True):
            name, url = (row[0] if len(row) > 0 else None), (row[1] if len(row) > 1 else None)
            if name and url:
                self._rows[_normalize(name)] = str(url).strip()

    def find(self, program: str, country: str) -> str:
        country_name = {"Argentina": "Argentina", "Republica Dominicana": "Dominicana"}.get(country, country)
        key = _normalize(f"{program} - {country_name}")
        if key in self._rows:
            return self._rows[key]
        raise KeyError(f"No se encontró ficha para {program} - {country}.")
