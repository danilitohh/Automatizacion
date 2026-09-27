"""Hoja de auditoría con los datos usados en cada intento de formulario."""

from __future__ import annotations

import re
from typing import Any

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


class FormDataReportService:
    """Agrega una vista por intento sin modificar las hojas de entrada."""

    SHEET_NAME = "Datos del formulario"
    COLUMNS = (
        "Hoja origen",
        "Fila Excel",
        "Intento",
        "URL formulario",
        "País",
        "Nivel",
        "Modalidad",
        "Tipo formulario",
        "Programa seleccionado",
        "Nombre",
        "Correo",
        "Teléfono",
        "Estado llenado",
        "Envío intentado",
        "Estado envío",
        "Resultado",
        "Error",
    )

    @staticmethod
    def _field_heading(field: dict[str, Any]) -> str:
        """Distingue controles con etiquetas iguales sin perder su nombre visible."""

        label = re.sub(r"\s+", " ", str(field.get("label") or "").strip())
        key = re.sub(r"\s+", " ", str(field.get("key") or "").strip())
        if not label:
            label = key or "Campo sin etiqueta"
        if key and key.casefold() != label.casefold():
            return f"Campo: {label} [{key}]"[:240]
        return f"Campo: {label}"[:240]

    @staticmethod
    def _safe_value(value: Any) -> str:
        """Mantiene identificadores como texto y evita fórmulas de origen web."""

        return str(value if value is not None else "")[:32767]

    @staticmethod
    def _form_url(row: dict[str, Any], result: dict[str, Any]) -> str:
        """Usa la URL realmente abierta cuando el Excel no incluía enlace."""

        for stage in result.get("stages") or []:
            if stage.get("stage") == "utel_open" and stage.get("url"):
                return str(stage["url"])
        return str(row.get("utel_url") or "")

    @staticmethod
    def _filled(attempt: dict[str, Any], result: dict[str, Any]) -> str:
        """No presenta datos generados como si ya se hubieran escrito en UTEL."""

        filled = attempt.get("form_filled")
        if filled is None:
            filled = any(
                stage.get("stage") == "utel_fill" and stage.get("status") == "PASS"
                for stage in result.get("stages") or []
            )
        return "Rellenado" if filled else "No confirmado"

    def append(self, workbook: Any, results: list[dict[str, Any]]) -> None:
        """Añade una fila por intento, con columnas para cada control capturado."""

        entries: list[tuple[dict[str, Any], dict[str, Any], dict[str, Any]]] = []
        field_headers: list[str] = []
        for item in results:
            row = item.get("row") or {}
            result = item.get("result") or {}
            attempts = result.get("form_attempts") or [result]
            for attempt in attempts:
                entries.append((row, result, attempt))
                for field in attempt.get("form_fields") or []:
                    heading = self._field_heading(field)
                    if heading not in field_headers:
                        field_headers.append(heading)

        sheet = workbook.create_sheet(self.SHEET_NAME)
        headers = [*self.COLUMNS, *field_headers]
        sheet.append(headers)
        for row, result, attempt in entries:
            attempted = attempt.get("utel_submission_attempted", result.get("utel_submission_attempted"))
            submission = "Sí" if attempted is True else "No" if attempted is False else "Sin confirmar"
            error = attempt.get("error") or result.get("error") or (
                result.get("summary") if result.get("status") == "FAIL" else ""
            )
            values = [
                row.get("sheet", ""),
                row.get("row_number", ""),
                attempt.get("attempt", 1),
                self._form_url(row, result),
                result.get("country") or row.get("country", ""),
                result.get("level") or row.get("level", ""),
                result.get("modality") or row.get("modality", ""),
                result.get("form_type") or row.get("form_type", ""),
                attempt.get("selected_program_name") or result.get("selected_program_name", ""),
                attempt.get("lead_name") or result.get("lead_name", ""),
                attempt.get("lead_email") or result.get("lead_email", ""),
                attempt.get("lead_phone") or result.get("lead_phone", ""),
                self._filled(attempt, result),
                submission,
                attempt.get("utel_submission") or result.get("utel_submission", ""),
                attempt.get("status") or result.get("status", ""),
                error,
            ]
            by_heading = {
                self._field_heading(field): field.get("value", "")
                for field in attempt.get("form_fields") or []
            }
            values.extend(by_heading.get(heading, "") for heading in field_headers)
            sheet.append(values)
            # Los teléfonos y demás cadenas se escriben explícitamente como
            # texto: Excel no debe recortar ceros ni ejecutar fórmulas externas.
            for cell in sheet[sheet.max_row]:
                if cell.column not in (2, 3):
                    cell.value = self._safe_value(cell.value)
                    cell.data_type = "s"

        # La hoja es una tabla de auditoría: cabecera fija, filtros y ancho
        # suficiente para revisar datos sin alterar el diseño original.
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, sheet.max_row)}"
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="17324D")
            cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
            cell.alignment = Alignment(vertical="center", horizontal="center")
        sheet.row_dimensions[1].height = 28
        widths = (20, 12, 10, 58, 20, 20, 18, 18, 36, 30, 38, 20, 18, 17, 18, 14, 60)
        for index in range(1, len(headers) + 1):
            sheet.column_dimensions[get_column_letter(index)].width = widths[index - 1] if index <= len(widths) else 30
        for cells in sheet.iter_rows(min_row=2):
            for cell in cells:
                cell.alignment = Alignment(vertical="center")
