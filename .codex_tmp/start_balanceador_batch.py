import json
import sys
from pathlib import Path

import httpx
from openpyxl import load_workbook


BASE_URL = "http://127.0.0.1:8010/api"
DEFAULT_INPUT_PATH = Path(r"C:\Users\danil\Downloads\Hoja de cálculo sin título.xlsx")


def pending_rows(input_path: Path) -> list[dict[str, object]]:
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    pending: list[dict[str, object]] = []
    for worksheet in workbook.worksheets:
        rows = list(worksheet.iter_rows(values_only=True))
        header_index = next(
            (
                index
                for index, row in enumerate(rows[:15])
                if "Programa" in row and "URL PAGE" in row and "URL_LEAD" in row
            ),
            None,
        )
        if header_index is None:
            continue
        headers = list(rows[header_index])
        program_index = headers.index("Programa")
        page_index = headers.index("URL PAGE")
        lead_index = headers.index("URL_LEAD")
        for row_number, row in enumerate(rows[header_index + 1 :], header_index + 2):
            program = row[program_index] if program_index < len(row) else None
            page_url = row[page_index] if page_index < len(row) else None
            lead_url = row[lead_index] if lead_index < len(row) else None
            if program and page_url and not str(lead_url or "").strip():
                pending.append({"sheet": worksheet.title, "row_number": row_number})
    return pending


def main() -> None:
    input_path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_INPUT_PATH
    selected_rows = pending_rows(input_path)
    if not selected_rows:
        raise RuntimeError("El archivo no contiene filas pendientes sin URL_LEAD.")
    config = {
        "name": "Bot nuevos productos - pendientes solo Balanceador",
        "environment": "sandbox",
        "dry_run": False,
        "country": "Argentina",
        "utel_url": "https://utel.edu.mx/argentina",
        "inconcert_url": "",
        "lead_origin_url": "",
        "lead_search_destination": "balanceador",
        "modality": "En linea",
        "level": "Licenciatura",
        "form_type": "tarjeta",
        "program_selection_strategy": "exact_match",
        "program_name": "",
        "submit_success_pattern": "Envío correcto|Pronto recibirás información",
        "submit_error_pattern": "Error al enviar|Contacta a soporte|error|invalido|inválido|obligatorio|requerido|fall",
        "browser": "chrome",
        "headless": False,
        "keep_browser_open": False,
        "lead": {
            "name": "pending",
            "email": "pending@testingUtel.com",
            "phone": "900000000",
        },
    }
    mapping = {
        "program_name": "Programa",
        "utel_url": "URL PAGE",
        "lead_url": "URL_LEAD",
        "selected_rows": selected_rows,
        "workflow_mode": "product_release",
    }

    with input_path.open("rb") as source:
        response = httpx.post(
            f"{BASE_URL}/bots/utel-inconcert/batch-run",
            files={
                "file": (
                    input_path.name,
                    source,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            data={"config": json.dumps(config), "mapping": json.dumps(mapping)},
            timeout=120,
        )
    response.raise_for_status()
    payload = response.json()
    print(json.dumps({"selected_rows": selected_rows}, ensure_ascii=False))
    print(json.dumps(payload, ensure_ascii=False))
    print(payload["job_id"])


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise
